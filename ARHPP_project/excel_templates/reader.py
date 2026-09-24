"""Read ARHPP Excel inputs into typed objects."""

from pathlib import Path
from typing import Optional
from openpyxl import load_workbook

from arhpp.core.types import (
    WellHeader, SurveyPoint, HoleSection, BHASection,
)


BASE = Path(__file__).resolve().parent.parent / "excel_input"


def _rows(path, sheet, skip=1):
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet]
    return list(ws.iter_rows(values_only=True))[skip:]


def read_well_header() -> WellHeader:
    r = _rows(BASE / "01_Well_Header.xlsx", "Well", 2)
    d = {row[0]: row[1] for row in r if row and row[0]}
    return WellHeader(
        well_name=str(d.get("Well Name", "WELL-1")),
        rig_name=str(d.get("Rig Name", "")),
        field=str(d.get("Field", "")),
        operator=str(d.get("Operator", "")),
        country=str(d.get("Country", "Kuwait")),
        rkb_ft=float(d.get("RKB") or 0),
        water_depth_ft=float(d.get("Water Depth") or 0),
        datum=str(d.get("Datum", "RKB")))


def read_survey():
    r = _rows(BASE / "02_Survey.xlsx", "Survey", 2)
    return [SurveyPoint(md=float(x[0]), inc=float(x[1]),
                         azi=float(x[2]))
            for x in r if x and x[0] is not None]


def read_hole_program():
    r = _rows(BASE / "03_Hole_Program.xlsx", "HoleProgram", 2)
    out = []
    for row in r:
        if not row or row[0] is None:
            continue
        out.append(HoleSection(
            name=str(row[0]), hole_id_in=float(row[1]),
            top_md=float(row[2]), bottom_md=float(row[3]),
            casing_od_in=float(row[4]) if row[4] else None,
            casing_id_in=float(row[5]) if row[5] else None,
            casing_shoe_md=float(row[6]) if row[6] else None))
    return out


def read_bha():
    r = _rows(BASE / "05_BHA_String.xlsx", "String", 2)
    out = []
    for row in r:
        if not row or row[0] is None:
            continue
        out.append(BHASection(
            name=str(row[1]), component_type=str(row[2]),
            od_in=float(row[3]), id_in=float(row[4]),
            length_ft=float(row[5])))
    return out


def read_mud_properties():
    r = _rows(BASE / "06_Mud_Properties.xlsx", "Mud", 2)
    out = {}
    for row in r:
        if not row or row[0] is None:
            continue
        out[str(row[0])] = {
            "mw": float(row[1]), "pv": float(row[2]),
            "yp": float(row[3]),
            "fann": [float(x) if x is not None else 0
                      for x in row[4:10]],
            "temp_f": float(row[10]) if row[10] else 100}
    return out


def read_fluid_segments_string():
    r = _rows(BASE / "07_Fluid_Segments_String.xlsx",
                "String_Fluids", 2)
    return [{"fluid_id": str(x[0]), "mw": float(x[1]),
             "top_md": float(x[2]), "bottom_md": float(x[3])}
            for x in r if x and x[0]]


def read_fluid_segments_annulus():
    r = _rows(BASE / "08_Fluid_Segments_Annulus.xlsx",
                "Annulus_Fluids", 2)
    return [{"fluid_id": str(x[0]), "mw": float(x[1]),
             "top_md": float(x[2]), "bottom_md": float(x[3])}
            for x in r if x and x[0]]


def read_bit_config():
    wb = load_workbook(BASE / "12_Bit_Config.xlsx", data_only=True)
    rows = list(wb["BitConfig"].iter_rows(values_only=True))
    bd = {r[0]: r[1] for r in rows[2:] if r and r[0]}
    rows = list(wb["Nozzles"].iter_rows(values_only=True))
    nozzles = [{"size_32nd_in": float(r[1]),
                "tfa_in2": float(r[2]) if r[2] else 0}
               for r in rows[2:] if r and r[0] is not None]
    return {"bit_diameter_in": float(bd.get("Bit Diameter", 8.5)),
            "cd": float(bd.get("Cd", 0.95)),
            "nozzles": nozzles}


def read_hydraulics_options():
    rows = list(load_workbook(BASE / "13_Hydraulics_Options.xlsx",
                                data_only=True)["Options"].iter_rows(
        values_only=True))
    d = {r[0]: r[1] for r in rows[2:] if r and r[0]}
    return {"q_gpm": float(d.get("Flow Rate", 650)),
            "sbp_psi": float(d.get("SBP", 250)),
            "rpm": float(d.get("RPM", 120)),
            "eccentricity": float(d.get("Eccentricity", 0.5))}


def read_utube_config():
    rows = list(load_workbook(BASE / "14_UTube_Config.xlsx",
                                data_only=True)["Config"].iter_rows(
        values_only=True))
    d = {r[0]: r[1] for r in rows[2:] if r and r[0]}
    b = lambda x: str(x).strip().upper() in ("TRUE", "1", "YES")
    return {"utube_enabled": b(d.get("U-Tube Enabled", "TRUE")),
            "choke_closed": b(d.get("Choke Closed", "FALSE"))}


def read_cuttings_config():
    rows = list(load_workbook(BASE / "17_Cuttings_Config.xlsx",
                                data_only=True)["Cuttings"].iter_rows(
        values_only=True))
    d = {r[0]: r[1] for r in rows[2:] if r and r[0]}
    return {"rop_ft_hr": float(d.get("ROP", 60)),
            "rpm": float(d.get("RPM", 120)),
            "rho_cut_ppg": float(d.get("Cuttings Density", 21.7)),
            "d_cut_in": float(d.get("Cuttings Size", 0.25))}


def read_gas_config():
    rows = list(load_workbook(BASE / "18_Gas_Config.xlsx",
                                data_only=True)["Gas"].iter_rows(
        values_only=True))
    d = {r[0]: r[1] for r in rows[2:] if r and r[0]}
    return {"gas_influx_bbl": float(d.get("Gas Influx", 0)),
            "gas_influx_md": float(d.get("Gas Influx MD", 0)),
            "gas_sg": float(d.get("Gas SG", 0.65)),
            "obm_dissolved_base": float(
                d.get("OBM Dissolved Base", 0.35)),
            "t_surface_f": float(d.get("Surface Temp", 90)),
            "geothermal_gradient": float(
                d.get("Geothermal Gradient", 0.015))}


def read_events_config():
    wb = load_workbook(BASE / "19_Events_Config.xlsx", data_only=True)

    def rd(s):
        rows = list(wb[s].iter_rows(values_only=True))
        return {r[0]: r[1] for r in rows[2:] if r and r[0]}

    def _b(x):
        return str(x).strip().upper() == "TRUE"

    k = rd("Kick"); nz = rd("Nozzle"); bl = rd("Balling")
    po = rd("PackOff"); wo = rd("Washout")

    return {
        "kick": {
            "q_in_gpm": float(k.get("Q_in", 650)),
            "q_out_gpm": float(k.get("Q_out", 650)),
            "pit_gain_bbl": float(k.get("Pit Gain", 0)),
            "connection_gas_units": float(k.get("Connection Gas", 0)),
            "background_gas_units": float(k.get("Background Gas", 0)),
            "trip_gas_units": float(k.get("Trip Gas", 0)),
            "pumps_off": _b(k.get("Pumps Off", "FALSE")),
            "tvd_ft": float(k.get("TVD", 0)),
            "mw_ppg": float(k.get("MW", 10)),
            "annular_fp_psi": float(k.get("Annular FP", 0)),
            "sbp_psi": float(k.get("SBP", 0)),
            "kick_duration_min": float(k.get("Kick Duration", 0))},
        "nozzle": {
            "expected_bit_dp_psi": float(nz.get("Expected Bit dP", 0)),
            "actual_bit_dp_psi": float(nz.get("Actual Bit dP", 0)),
            "tfa_expected_in2": float(nz.get("TFA Expected", 0)),
            "cd": float(nz.get("Cd", 0.95))},
        "balling": {
            "spp_baseline_psi": float(bl.get("SPP Baseline", 0)),
            "spp_current_psi": float(bl.get("SPP Current", 0)),
            "torque_baseline_ftlb": float(bl.get("Torque Baseline", 0)),
            "torque_current_ftlb": float(bl.get("Torque Current", 0)),
            "rop_baseline_ft_hr": float(bl.get("ROP Baseline", 0)),
            "rop_current_ft_hr": float(bl.get("ROP Current", 0)),
            "wob_baseline_klb": float(bl.get("WOB Baseline", 0)),
            "wob_current_klb": float(bl.get("WOB Current", 0)),
            "formation": str(bl.get("Formation", "shale")),
            "mud_type": str(bl.get("Mud Type", "OBM"))},
        "packoff": {
            "spp_baseline_psi": float(po.get("SPP Baseline", 0)),
            "spp_current_psi": float(po.get("SPP Current", 0)),
            "torque_baseline_ftlb": float(po.get("Torque Baseline", 0)),
            "torque_current_ftlb": float(po.get("Torque Current", 0)),
            "drag_baseline_klb": float(po.get("Drag Baseline", 0)),
            "drag_current_klb": float(po.get("Drag Current", 0)),
            "flow_in_gpm": float(po.get("Flow In", 0)),
            "flow_out_gpm": float(po.get("Flow Out", 0))},
        "washout": {
            "spp_baseline_psi": float(wo.get("SPP Baseline", 0)),
            "spp_current_psi": float(wo.get("SPP Current", 0)),
            "spp_noise_std_psi": float(wo.get("SPP Noise Std", 0)),
            "flow_in_gpm": float(wo.get("Flow In", 0)),
            "flow_out_gpm": float(wo.get("Flow Out", 0)),
            "q_step_gpm": float(wo.get("Q Step", 0)),
            "spp_response_ratio": float(
                wo.get("SPP Response Ratio", 1.0))},
    }
