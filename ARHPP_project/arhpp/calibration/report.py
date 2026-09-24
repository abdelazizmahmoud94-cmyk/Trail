"""Calibration report writer."""

from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from arhpp.calibration.params import REGISTRY

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")


def _style_header(ws, row=1):
    for c in ws[row]:
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")


def write_report(result, path) -> Path:
    p = Path(path)
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Summary")
    ws["A1"] = "ARHPP Calibration Report"
    ws["A1"].font = TITLE_FONT
    rows = [
        ("Initial Score", round(result.initial_score, 5)),
        ("Final Score", round(result.best_score, 5)),
        ("Improvement %", round(result.improvement_pct, 2)),
        ("Iterations", result.n_iterations),
        ("Wells Calibrated", len(result.well_targets)),
    ]
    r = 3
    for k, v in rows:
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r, column=2, value=v)
        r += 1

    ws = wb.create_sheet("Parameters")
    headers = ["Param", "Default", "Calibrated", "Delta", "Lo", "Hi",
                "Unit", "Category", "Description"]
    for j, h in enumerate(headers, 1):
        ws.cell(row=1, column=j, value=h)
    _style_header(ws)
    for i, spec in enumerate(REGISTRY, 2):
        ws.cell(row=i, column=1, value=spec.name)
        ws.cell(row=i, column=2, value=spec.default)
        new_v = result.best_params.get(spec.name, spec.default) if hasattr(result, "best_params") else result.best_profile.values.get(spec.name, spec.default)
        ws.cell(row=i, column=3, value=round(new_v, 5))
        ws.cell(row=i, column=4, value=round(new_v - spec.default, 5))
        ws.cell(row=i, column=5, value=spec.lo)
        ws.cell(row=i, column=6, value=spec.hi)
        ws.cell(row=i, column=7, value=spec.unit)
        ws.cell(row=i, column=8, value=spec.category)
        ws.cell(row=i, column=9, value=spec.description)

    ws = wb.create_sheet("PerWell")
    headers = ["Well", "Composite", "n_obj",
                "dBHP (psi)", "dECD (ppg)", "dSPP (psi)", "dPP (ppg)"]
    for j, h in enumerate(headers, 1):
        ws.cell(row=1, column=j, value=h)
    _style_header(ws)
    for i, t in enumerate(result.well_targets, 2):
        ws.cell(row=i, column=1, value=t.well_id)
        ws.cell(row=i, column=2, value=round(t.composite, 5))
        ws.cell(row=i, column=3, value=t.n_objectives)
        ws.cell(row=i, column=4, value=round(t.e_bhp_psi, 2))
        ws.cell(row=i, column=5, value=round(t.e_ecd_ppg, 4))
        ws.cell(row=i, column=6, value=round(t.e_spp_psi, 2))
        ws.cell(row=i, column=7, value=round(t.e_pp_ppg, 4))

    ws = wb.create_sheet("Convergence")
    ws.cell(row=1, column=1, value="Iteration")
    ws.cell(row=1, column=2, value="Score")
    _style_header(ws)
    for i, (it, score, _) in enumerate(result.history, 2):
        ws.cell(row=i, column=1, value=it)
        ws.cell(row=i, column=2, value=round(score, 6))

    ws = wb.create_sheet("Stages")
    ws.cell(row=1, column=1, value="Stage Description")
    _style_header(ws)
    for i, s in enumerate(result.stages, 2):
        ws.cell(row=i, column=1, value=s)

    for sh in wb.sheetnames:
        w = wb[sh]
        for col in w.columns:
            try:
                m = max(len(str(c.value or "")) for c in col) + 3
                w.column_dimensions[col[0].column_letter].width = min(m, 40)
            except Exception:
                pass

    wb.save(p)
    return p
