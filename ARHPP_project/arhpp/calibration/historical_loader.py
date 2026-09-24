"""Load historical well data from Excel / CSV."""

from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional

import csv

from openpyxl import load_workbook


@dataclass
class HistoricalWell:
    well_id: str
    md_ft: float = 0.0
    tvd_ft: float = 0.0
    hole_id_in: float = 8.5
    pipe_od_in: float = 5.0
    pipe_id_in: float = 4.276
    mw_in_ppg: float = 10.0
    mw_out_ppg: float = 10.0
    temp_in_f: float = 100.0
    temp_out_f: float = 150.0
    bht_f: float = 250.0
    fann_600: float = 60.0
    fann_300: float = 35.0
    fann_200: float = 27.0
    fann_100: float = 18.0
    fann_6: float = 6.0
    fann_3: float = 4.0
    gels_10s: float = 6.0
    gels_10min: float = 15.0
    q_gpm: float = 600.0
    rpm: float = 120.0
    rop_ft_hr: float = 40.0
    wob_klb: float = 25.0
    sbp_psi: float = 0.0
    eccentricity: float = 0.5
    pwd_bhp_psi: Optional[float] = None
    pwd_ecd_ppg: Optional[float] = None
    spp_measured_psi: Optional[float] = None
    q_out_measured_gpm: Optional[float] = None
    had_kick: bool = False
    had_losses: bool = False
    kick_severity: str = "none"
    loss_class: str = "none"
    pp_reference_ppg: Optional[float] = None
    fg_reference_ppg: Optional[float] = None
    formation: str = ""
    tfa_in2: float = 0.45
    cd: float = 0.95


def _f(v, d=0.0):
    try:
        return float(v) if v not in (None, "") else d
    except (TypeError, ValueError):
        return d


def _b(v) -> bool:
    return str(v).strip().upper() in ("TRUE", "1", "YES", "Y")


def _opt_f(v):
    try:
        if v in (None, ""):
            return None
        return float(v)
    except (TypeError, ValueError):
        return None


def load_well_from_row(row: dict, well_id: str) -> HistoricalWell:
    return HistoricalWell(
        well_id=well_id,
        md_ft=_f(row.get("MD_ft")),
        tvd_ft=_f(row.get("TVD_ft")),
        hole_id_in=_f(row.get("Hole_ID_in"), 8.5),
        pipe_od_in=_f(row.get("Pipe_OD_in"), 5.0),
        pipe_id_in=_f(row.get("Pipe_ID_in"), 4.276),
        mw_in_ppg=_f(row.get("MW_in_ppg"), 10.0),
        mw_out_ppg=_f(row.get("MW_out_ppg"), 10.0),
        temp_in_f=_f(row.get("Temp_in_F"), 100.0),
        temp_out_f=_f(row.get("Temp_out_F"), 150.0),
        bht_f=_f(row.get("BHT_F"), 250.0),
        fann_600=_f(row.get("Fann600"), 60.0),
        fann_300=_f(row.get("Fann300"), 35.0),
        fann_200=_f(row.get("Fann200"), 27.0),
        fann_100=_f(row.get("Fann100"), 18.0),
        fann_6=_f(row.get("Fann6"), 6.0),
        fann_3=_f(row.get("Fann3"), 4.0),
        gels_10s=_f(row.get("Gels_10s"), 6.0),
        gels_10min=_f(row.get("Gels_10min"), 15.0),
        q_gpm=_f(row.get("Q_gpm"), 600.0),
        rpm=_f(row.get("RPM"), 120.0),
        rop_ft_hr=_f(row.get("ROP_ft_hr"), 40.0),
        wob_klb=_f(row.get("WOB_klb"), 25.0),
        sbp_psi=_f(row.get("SBP_psi"), 0.0),
        eccentricity=_f(row.get("Eccentricity"), 0.5),
        pwd_bhp_psi=_opt_f(row.get("PWD_BHP_psi")),
        pwd_ecd_ppg=_opt_f(row.get("PWD_ECD_ppg")),
        spp_measured_psi=_opt_f(row.get("SPP_psi")),
        q_out_measured_gpm=_opt_f(row.get("Q_out_gpm")),
        had_kick=_b(row.get("HadKick", "FALSE")),
        had_losses=_b(row.get("HadLosses", "FALSE")),
        kick_severity=str(row.get("KickSeverity", "none") or "none"),
        loss_class=str(row.get("LossClass", "none") or "none"),
        pp_reference_ppg=_opt_f(row.get("PP_ref_ppg")),
        fg_reference_ppg=_opt_f(row.get("FG_ref_ppg")),
        formation=str(row.get("Formation", "") or ""),
        tfa_in2=_f(row.get("TFA_in2"), 0.45),
        cd=_f(row.get("Cd"), 0.95),
    )


def load_wells_from_excel(path) -> List[HistoricalWell]:
    wb = load_workbook(path, data_only=True)
    if "Wells" not in wb.sheetnames:
        raise ValueError(f"No 'Wells' sheet in {path}")
    ws = wb["Wells"]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h) if h is not None else "" for h in rows[0]]
    wells = []
    for i, r in enumerate(rows[1:], start=2):
        if not r or all(v in (None, "") for v in r):
            continue
        row = dict(zip(headers, r))
        wid = str(row.get("Well_ID") or f"WELL_{i}")
        wells.append(load_well_from_row(row, wid))
    return wells


def load_wells_from_csv_dir(dir_path) -> List[HistoricalWell]:
    d = Path(dir_path)
    wells = []
    for f in sorted(d.glob("*.csv")):
        with open(f, "r", encoding="utf-8") as fh:
            reader = csv.DictReader(fh)
            row = next(reader, None)
            if row:
                wells.append(load_well_from_row(row, f.stem))
    return wells
