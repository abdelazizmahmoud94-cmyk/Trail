"""ARHPP P8 - Full end-to-end run."""

from pathlib import Path
from openpyxl import load_workbook

from arhpp.geometry.survey import compute_tvd, md_to_tvd
from arhpp.geometry.bha import (
    stack_bha, string_id_at_md, string_od_at_md,
)
from arhpp.geometry.well_geometry import effective_hole_id
from arhpp.fluids.rheology import from_fann_6
from arhpp.fluids.fluid_tracking import (
    build_string_column, build_annulus_column,
)
from arhpp.hydraulics.bit import Nozzle
from arhpp.hydraulics.ledger import build_ledger
from arhpp.cuttings.transport import compute_cuttings_transport
from arhpp.gas.behaviour import compute_gas_behaviour
from arhpp.events.kick import KickInputs, compute_kick
from arhpp.events.nozzle_plug import compute_nozzle_plug
from arhpp.events.bit_balling import (
    BitBallingInputs, compute_bit_balling,
)
from arhpp.events.packoff import PackOffInputs, compute_packoff
from arhpp.events.washout import WashoutInputs, compute_washout
from arhpp.pore_pressure.kuwait_calibration import (
    pp_ppg_at_md, fg_ppg_at_md,
)
from arhpp.meta.diagnostics import (
    build_diagnostics, diagnostics_to_rows,
)
from arhpp.meta.dashboard import (
    DashboardInputs, build_dashboard, dashboard_to_rows,
)
from excel_templates import reader as xl


BASE = Path(__file__).resolve().parent.parent
IN = BASE / "excel_input"


def write_diag(d):
    p = IN / "A4_Output_Diagnostics.xlsx"
    if not p.exists():
        return
    wb = load_workbook(p)
    ws = wb["Ledger"]
    for i, r in enumerate(diagnostics_to_rows(d), 3):
        for j, k in enumerate(["Component", "Value_psi", "Sign",
                                "Category", "Fraction_of_BHP",
                                "Comment"], 1):
            ws.cell(row=i, column=j, value=r[k])
    wb.save(p)


def write_dash(d):
    p = IN / "A6_Dashboard.xlsx"
    if not p.exists():
        return
    wb = load_workbook(p)
    ws = wb["Dashboard"]
    for i, r in enumerate(dashboard_to_rows(d), 3):
        ws.cell(row=i, column=1, value=r["Item"])
        ws.cell(row=i, column=2, value=r["Value"])
    ws = wb["Alerts"]
    for i, a in enumerate(d.alerts, 3):
        ws.cell(row=i, column=1, value=a)
    wb.save(p)


def main():
    print("ARHPP - P8 Full Run\n" + "=" * 65)

    # Load
    survey = compute_tvd(xl.read_survey())
    sections = xl.read_hole_program()
    bha = stack_bha(xl.read_bha())
    mud = xl.read_mud_properties()
    raw_s = xl.read_fluid_segments_string()
    raw_a = xl.read_fluid_segments_annulus()
    opts = xl.read_hydraulics_options()
    bit_cfg = xl.read_bit_config()
    utube_cfg = xl.read_utube_config()
    cut_cfg = xl.read_cuttings_config()
    gas_cfg = xl.read_gas_config()
    ev_cfg = xl.read_events_config()

    def enrich(rows):
        out = []
        for r in rows:
            m = mud[r["fluid_id"]]
            hb = from_fann_6(*m["fann"])["herschel_bulkley"]
            out.append({**r, "tau_y": hb.tau_y,
                        "k": hb.k, "n": hb.n})
        return out

    raw_s = enrich(raw_s)
    raw_a = enrich(raw_a)

    string_col = build_string_column(
        raw_s, lambda md: string_id_at_md(bha, md), survey)
    ann_col = build_annulus_column(
        raw_a,
        lambda md: effective_hole_id(sections, md),
        lambda md: string_od_at_md(bha, md),
        survey)

    for seg, r in zip(string_col.segments, raw_s):
        seg.tau_y, seg.k, seg.n = r["tau_y"], r["k"], r["n"]
    for seg, r in zip(ann_col.segments, raw_a):
        seg.tau_y, seg.k, seg.n = r["tau_y"], r["k"], r["n"]

    nozzles = [Nozzle(size_32nd_in=n["size_32nd_in"],
                       tfa_in2=n["tfa_in2"])
                for n in bit_cfg["nozzles"]]

    q_out = ev_cfg["kick"]["q_out_gpm"]
    bundle = build_ledger(
        q_pump_gpm=opts["q_gpm"], sbp_psi=opts["sbp_psi"],
        string_col=string_col, annulus_col=ann_col,
        bha_sections=bha, hole_sections=sections, survey=survey,
        nozzles=nozzles,
        bit_diameter_in=bit_cfg["bit_diameter_in"],
        cd=bit_cfg["cd"],
        eccentricity=opts["eccentricity"], rpm=opts["rpm"],
        q_out_measured_gpm=q_out,
        utube_enabled=utube_cfg["utube_enabled"],
        choke_closed=utube_cfg["choke_closed"])
    L = bundle["ledger"]

    cut_res = compute_cuttings_transport(
        cut_cfg["rop_ft_hr"], cut_cfg["rpm"], opts["q_gpm"],
        sections, bha, ann_col, survey,
        d_cut_in=cut_cfg["d_cut_in"],
        rho_cut_ppg=cut_cfg["rho_cut_ppg"])
    L.annular_friction += cut_res.total_extra_dp_psi
    L.bhp += cut_res.total_extra_dp_psi

    gas_res = compute_gas_behaviour(
        gas_cfg["gas_influx_bbl"], gas_cfg["gas_influx_md"],
        string_col, ann_col, survey,
        gas_sg=gas_cfg["gas_sg"],
        t_surface_f=gas_cfg["t_surface_f"],
        temp_gradient_f_per_ft=gas_cfg["geothermal_gradient"],
        obm_dissolved_base=gas_cfg["obm_dissolved_base"])
    L.gas_effect = gas_res.total_bhp_reduction_psi
    L.bhp += L.gas_effect
    L.ecd = L.bhp / max(0.052 * L.tvd_ref, 1e-6)

    kick = compute_kick(
        KickInputs(**ev_cfg["kick"]),
        ev_cfg["kick"]["kick_duration_min"])
    nz = compute_nozzle_plug(
        ev_cfg["nozzle"]["expected_bit_dp_psi"],
        ev_cfg["nozzle"]["actual_bit_dp_psi"],
        ev_cfg["kick"]["q_in_gpm"],
        ev_cfg["kick"]["mw_ppg"],
        ev_cfg["nozzle"]["tfa_expected_in2"],
        ev_cfg["nozzle"]["cd"])
    bl = compute_bit_balling(BitBallingInputs(**ev_cfg["balling"]))
    po = compute_packoff(PackOffInputs(**ev_cfg["packoff"]))
    wo = compute_washout(WashoutInputs(**ev_cfg["washout"]))

    TD = sections[-1].bottom_md
    tvd_td = md_to_tvd(survey, TD)
    pp = pp_ppg_at_md(TD, tvd_td)
    fg = fg_ppg_at_md(TD, tvd_td)
    mw_td = ann_col.segments[-1].mw if ann_col.segments else 10.0

    diag = build_diagnostics(L, spp_psi=bundle["spp_psi"])

    dash = build_dashboard(DashboardInputs(
        ledger=L,
        q_pump_gpm=opts["q_gpm"], q_out_gpm=q_out,
        q_effective_gpm=bundle["q_effective_gpm"],
        q_utube_gpm=bundle["utube"].q_utube_gpm,
        q_loss_gpm=bundle["losses"].q_loss_gpm,
        spp_model_psi=bundle["spp_psi"],
        spp_measured_psi=ev_cfg["balling"]["spp_current_psi"],
        pp_ppg=pp, fg_ppg=fg, mw_ppg=mw_td,
        kick_severity=kick.severity,
        kick_probability=kick.probability,
        loss_class=bundle["losses"].loss_class,
        losses_fraction=bundle["losses"].loss_fraction,
        packoff_risk=po.risk, washout_risk=wo.risk,
        nozzle_risk=nz.risk,
        nozzle_plugging_pct=nz.plugging_percent,
        balling_risk=bl.risk,
        utube_severity=bundle["utube"].severity_index,
        utube_direction=bundle["utube"].flow_direction,
        cuttings_extra_ecd_ppg=cut_res.total_extra_ecd_ppg,
        gas_bhp_reduction_psi=gas_res.total_bhp_reduction_psi,
        sensor_agreement=0.95, model_agreement=0.90))

    print(f"\nTD = {TD:.0f} ft MD / {tvd_td:.0f} ft TVD")
    print(f"MW={mw_td:.2f}  PP={pp:.2f}  FG={fg:.2f}")
    print(f"\nBHP={L.bhp:.2f} psi  ECD={L.ecd:.4f} ppg  "
          f"ESD={dash.esd_ppg:.4f} ppg")
    print(f"Q_eff={bundle['q_effective_gpm']:.1f}  "
          f"Q_utube={bundle['utube'].q_utube_gpm:+.2f}  "
          f"Q_loss={bundle['losses'].q_loss_gpm:.2f}")
    print(f"Kick={kick.severity.upper()} "
          f"({kick.probability*100:.0f}%)  "
          f"Loss={bundle['losses'].loss_class.upper()}")
    print(f"Confidence="
          f"{dash.confidence.overall_confidence*100:.0f}%  "
          f"Overall={dash.overall_risk}")
    if dash.alerts:
        print("Alerts:", " | ".join(dash.alerts))

    write_diag(diag)
    write_dash(dash)
    print("\nWrote A4_Output_Diagnostics.xlsx / A6_Dashboard.xlsx")


